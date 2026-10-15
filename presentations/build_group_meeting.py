#!/usr/bin/env python3
"""Build the editable 16:9 project update deck from verified local results."""
import csv
import json
import math
import os
import shutil
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE=Path(__file__).resolve().parent
ROOT=HERE.parent
VIS=ROOT.parent/'ht_cupy_bridge_ws'/'results'/'xytheta_live_planning_2026-10-12'
NAVY=RGBColor(13,29,49); INK=RGBColor(25,42,62); TEAL=RGBColor(0,151,151)
BLUE=RGBColor(38,116,187); ORANGE=RGBColor(225,128,45); RED=RGBColor(188,76,65)
PALE=RGBColor(242,247,249); MID=RGBColor(91,111,129); WHITE=RGBColor(255,255,255)
FONT='Noto Sans CJK SC'
comparison_file=Path(os.environ.get('HT_PPT_COMPARISON_JSON',HERE/'latest_current_vs_upstream.json'))
comparison=json.loads(comparison_file.read_text()) if comparison_file.exists() else None
fresh=bool(comparison and comparison.get('status')=='complete')
RUN_DATE=(Path(comparison['run_dir']).name.split('_')[0] if fresh else '2026-10-12')
OUT=HERE/f'HT_TARE_XYtheta_组会汇报_{RUN_DATE}.pptx'
LATEST_OUT=HERE/'HT_TARE_XYtheta_组会汇报_最新.pptx'

prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
blank=prs.slide_layouts[6]

def txt(slide,x,y,w,h,s,size=20,color=INK,bold=False,align=None):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.02)
    tf.margin_top=tf.margin_bottom=Inches(.01)
    p=tf.paragraphs[0];p.text=s;p.font.name=FONT;p.font.size=Pt(size)
    p.font.bold=bold;p.font.color.rgb=color
    if align is not None:p.alignment=align
    return box

def rect(slide,x,y,w,h,fill=PALE,line=None,radius=False):
    shp=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x),Inches(y),Inches(w),Inches(h))
    shp.fill.solid();shp.fill.fore_color.rgb=fill
    shp.line.fill.background() if line is None else None
    if line is not None:shp.line.color.rgb=line
    return shp

def slide(title,kicker=None):
    s=prs.slides.add_slide(blank)
    s.background.fill.solid();s.background.fill.fore_color.rgb=WHITE
    rect(s,0,0,13.333,.11,TEAL)
    txt(s,.62,.3,11.9,.57,title,27,NAVY,True)
    if kicker:txt(s,.65,.92,11.8,.35,kicker,11,MID)
    txt(s,.66,7.13,9,.2,f'HT × TARE × XYθ  |  CMU Garage 仿真  |  {RUN_DATE.replace("-", ".")}',9,MID)
    txt(s,12.22,7.10,.45,.25,f'{len(prs.slides):02d}',10,MID,align=PP_ALIGN.RIGHT)
    return s

def card(s,x,y,w,h,head,body,accent=TEAL):
    rect(s,x,y,w,h,PALE,radius=True)
    rect(s,x,y,.08,h,accent)
    txt(s,x+.24,y+.19,w-.46,.36,head,17,INK,True)
    txt(s,x+.24,y+.69,w-.46,h-.85,body,14,MID)

def picture(s,path,x,y,w,h):
    from PIL import Image
    with Image.open(path) as im:iw,ih=im.size
    scale=min(w/iw,h/ih);pw=iw*scale;ph=ih*scale
    s.shapes.add_picture(str(path),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),
                         width=Inches(pw),height=Inches(ph))

def line(s,x1,y1,x2,y2,color=TEAL,width=2):
    c=s.shapes.add_connector(1,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    c.line.color.rgb=color;c.line.width=Pt(width)

# 1. Title
s=prs.slides.add_slide(blank);s.background.fill.solid();s.background.fill.fore_color.rgb=NAVY
rect(s,.0,.0,.14,7.5,TEAL)
txt(s,.86,.76,10.9,.4,f'项目阶段汇报  ·  {RUN_DATE.replace("-", ".")}',17,RGBColor(126,212,211),True)
txt(s,.85,1.58,11.7,1.45,'面向地面机器人探索的\n方向感知 XYθ 局部规划',36,WHITE,True)
rect(s,.9,3.55,11.5,.025,RGBColor(74,104,130))
txt(s,.9,3.94,10.8,.8,'TARE 全局探索引导  +  HT 八方向可通过性  +  Hybrid A* 差速车路径',19,WHITE)
txt(s,.9,6.58,10.6,.36,'阶段：Garage 闭环仿真原型；未完成实车验证',14,RGBColor(186,207,221))

# 2. Motivation
s=slide('核心问题：地形可通过性取决于车头方向','为什么仅在 XY 平面判断“这个格子能不能走”不够')
card(s,.65,1.60,3.78,3.28,'同一位置，方向不同','上坡、横坡、台阶边缘等地形，前进方向变化会改变车体支撑与行驶风险。',TEAL)
card(s,4.77,1.60,3.78,3.28,'探索不能只找高收益点','目标视点即使有覆盖收益，也必须存在从当前姿态到目标的可执行路线。',BLUE)
card(s,8.90,1.60,3.78,3.28,'差速车还要考虑转向','车辆可前进、倒车、原地旋转；路径的姿态序列比 XY 折线多一维。',ORANGE)
rect(s,.65,5.29,12.03,1.05,NAVY,radius=True)
txt(s,.95,5.52,11.45,.55,'目标：在 P(x,y,θ) 上拒绝不安全运动边，再选择可到达的探索视点。',21,WHITE,True)

# 3. Architecture
s=slide('系统链路：感知 → 方向地图 → 探索 → 执行','Garage 闭环试验中的数据流；TARE、CuPy 高程图与仿真处于独立工作空间')
boxes=[('雷达点云 / 里程计','CMU 仿真输入'),('CuPy 高程图','构建局部地形'),('HT 八方向地图','P(x,y,θₖ)'),
       ('TARE 探索视点','覆盖收益 / 全局引导'),('XYθ Hybrid A*','硬约束路线'),('差速跟踪器','/cmd_vel')]
xs=[.6,2.73,4.86,6.99,9.12,11.25]
for i,(a,b) in enumerate(boxes):
    rect(s,xs[i],2.18,1.56,1.55,PALE,radius=True)
    txt(s,xs[i]+.09,2.43,1.38,.56,a,16,INK,True,PP_ALIGN.CENTER)
    txt(s,xs[i]+.09,3.13,1.38,.3,b,10,MID,align=PP_ALIGN.CENTER)
    if i<5:txt(s,xs[i]+1.61,2.72,.43,.5,'→',22,TEAL,True,PP_ALIGN.CENTER)
card(s,.76,4.42,5.75,1.48,'规划输入','HT 的 8 张方向概率图 + 当前 (x,y,θ) + 候选探索视点',TEAL)
card(s,6.80,4.42,5.75,1.48,'规划输出','/ht_xytheta_path；跟踪器以 20 Hz 发布速度指令',BLUE)

# 4. XYtheta representation
s=slide('XYθ 空间：八张方向栅格构成三维地形场','同一 XY 位置在 θ=0°、45°、…、315° 上有不同的 HT 预测概率')
picture(s,VIS/'xytheta_3d_keyframe.png',.53,1.30,7.44,5.56)
card(s,8.18,1.61,4.47,1.25,'HT 地图','8 个方向层；原始栅格约 0.041 m',TEAL)
card(s,8.18,3.07,4.47,1.25,'规划状态','搜索键：XY 0.3 m、16 个朝向格',BLUE)
card(s,8.18,4.53,4.47,1.25,'颜色含义','绿：概率高；红：概率低；灰：未观测',ORANGE)
txt(s,8.25,6.15,4.16,.35,'示意来自实际仿真 HT 地图与规划路线',11,MID)

# 5. Planner
s=slide('局部路径：从当前姿态搜索可执行轨迹','Hybrid A* 保留连续姿态；每段轨迹都做地形与几何检查')
card(s,.62,1.52,3.79,3.34,'状态与动作','起点：(x, y, θ)\n动作：前进 / 倒车圆弧、直行、原地旋转\n搜索目标：候选视点位置',TEAL)
card(s,4.77,1.52,3.79,3.34,'硬约束','沿运动边按 ≤0.1 m / 0.12 rad 采样\n车体占地、HT 方向概率及几何碰撞均合格才保留该边',BLUE)
card(s,8.90,1.52,3.79,3.34,'时间预算','单候选搜索上限 0.15 s / 8000 展开\n前进 0.5 m/s；倒车 0.25 m/s；角速度 ≤0.8 rad/s',ORANGE)
rect(s,.65,5.27,12.03,1.12,NAVY,radius=True)
txt(s,.95,5.48,11.38,.68,'现阶段：TARE 负责视点采样与收益；XYθ 负责可达性和局部路线。',20,WHITE,True)

# 6. Live visualization
s=slide('实测规划：八个方向层中的同一条路线','Garage 探索过程中的一次规划：车约 (0.37, 5.25) m、103° → 视点约 (3.19, 7.47) m')
picture(s,VIS/'xytheta_grid_keyframe.png',.44,1.31,12.45,5.50)
txt(s,.66,6.86,11.7,.25,'黑虚线：整条路径的 XY 投影；青色：该 θ 层的前进路段；蓝星：目标视点。',11,MID)

# 7. Results
if fresh:
    s=slide('最新 Garage 实验：路径访问范围随时间变化',
            f'原版与当前 XYθ 版各 {comparison["duration_s"]} 秒；到访栅格是车辆轨迹经过的不同 1 m XY 格')
    run_dir=Path(comparison['run_dir'])
    curves=[];last_current_motion=0
    for name,color in [('original',BLUE),('current',TEAL)]:
        seen=set();points=[];previous_distance=None
        with (run_dir/name/'timeseries.csv').open(newline='') as f:
            for row in csv.DictReader(f):
                if row['x_m'] and row['y_m']:
                    seen.add((math.floor(float(row['x_m'])),math.floor(float(row['y_m']))))
                t=float(row['t_s'])
                distance=float(row['distance_m'])
                if name=='current' and previous_distance is not None and distance-previous_distance>.01:
                    last_current_motion=t
                previous_distance=distance
                if not points or t-points[-1][0]>=5 or t>=comparison['duration_s']-.6:
                    points.append((t,len(seen)))
        curves.append((name,color,points))
    ymax=max(10,max(v for _,_,pts in curves for _,v in pts))
    ymax=math.ceil(ymax/20)*20
    rect(s,.66,1.50,7.42,4.79,PALE,radius=True)
    txt(s,.94,1.73,6.3,.35,'到访 1 m XY 栅格数',17,INK,True)
    ox,oy,w,h=1.13,5.64,6.35,3.19
    line(s,ox,oy,ox+w,oy,MID,1.4);line(s,ox,oy,ox,oy-h,MID,1.4)
    for value in (0,ymax//2,ymax):
        yy=oy-h*value/ymax
        line(s,ox,yy,ox+w,yy,RGBColor(215,226,231),.6)
        txt(s,.73,yy-.12,.34,.24,str(value),9,MID,align=PP_ALIGN.RIGHT)
    for t in (0,comparison['duration_s']//2,comparison['duration_s']):
        x=ox+w*t/comparison['duration_s']
        txt(s,x-.25,oy+.12,.50,.24,str(t),10,MID,align=PP_ALIGN.CENTER)
    for name,color,points in curves:
        for (t1,v1),(t2,v2) in zip(points,points[1:]):
            line(s,ox+w*t1/comparison['duration_s'],oy-h*v1/ymax,
                 ox+w*t2/comparison['duration_s'],oy-h*v2/ymax,color,2.6)
    txt(s,1.40,6.05,2.45,.28,'— 原版 TARE',12,BLUE,True)
    txt(s,3.90,6.05,2.50,.28,'— 当前 XYθ',12,TEAL,True)
    card(s,8.43,1.51,4.19,1.16,f'{comparison["original"]["distance_m"]:.2f} m','原版行驶距离',BLUE)
    card(s,8.43,2.83,4.19,1.16,f'{comparison["current"]["distance_m"]:.2f} m','当前 XYθ 行驶距离',TEAL)
    card(s,8.43,4.15,4.19,1.16,
         f'{comparison["original"]["visited_1m_xy_cells"]} / {comparison["current"]["visited_1m_xy_cells"]} 格',
         '原版 / 当前的轨迹到访栅格',ORANGE)
    rect(s,.67,6.43,11.98,.39,RGBColor(255,240,224),radius=True)
    plateau=(comparison['duration_s']-last_current_motion)>15
    txt(s,.86,6.47,11.58,.27,
        f'当前版约 {last_current_motion:.0f} 秒后里程未再增加；轨迹到访格不是探索体积，单次实验不能归因。' if plateau else
        '这是单次端到端试验；轨迹到访格不是探索体积，也不能单独归因于 HT 或 XYθ。',12,RED,True)
else:
    s=slide('Garage 闭环：能够推进，但尚未完成探索','90 秒记录取自许可超时修正前的运行；体积为新增的、被有效光线穿过的自由体素')
    data=[(0,0,0),(30,240.4,2.0),(60,376.5,4.1),(90,462.6,5.5)]
    rect(s,.66,1.50,7.42,4.79,PALE,radius=True)
    txt(s,.94,1.73,6.3,.35,'新增自由体积（m³）',17,INK,True)
    ox,oy,w,h=1.13,5.64,6.35,3.19
    line(s,ox,oy,ox+w,oy,MID,1.4);line(s,ox,oy,ox,oy-h,MID,1.4)
    for value in (0,250,500):
        yy=oy-h*value/500
        line(s,ox,yy,ox+w,yy,RGBColor(215,226,231),.6)
        txt(s,.76,yy-.12,.30,.24,str(value),9,MID,align=PP_ALIGN.RIGHT)
    for t,v,_ in data:
        x=ox+w*t/90;y=oy-h*v/500
        if t>0:
            prev=data[data.index((t,v,_))-1];px=ox+w*prev[0]/90;py=oy-h*prev[1]/500
            line(s,px,py,x,y,TEAL,3)
        circ=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x-.06),Inches(y-.06),Inches(.12),Inches(.12))
        circ.fill.solid();circ.fill.fore_color.rgb=TEAL;circ.line.fill.background()
        txt(s,x-.22,oy+.12,.46,.25,str(t),10,MID,align=PP_ALIGN.CENTER)
    txt(s,7.02,5.84,.65,.24,'时间 / s',10,MID)
    card(s,8.43,1.51,4.19,1.16,'+462.6 m³','90 秒内新增自由体积',TEAL)
    card(s,8.43,2.83,4.19,1.16,'5.51 m','同期里程计行驶距离',BLUE)
    card(s,8.43,4.15,4.19,1.16,'81 / 82 帧','有效注册扫描 / 总扫描',ORANGE)
    rect(s,.67,6.43,11.98,.39,RGBColor(255,240,224),radius=True)
    txt(s,.86,6.47,11.58,.27,'90 秒时仍在探索；其后当前位置被 HT 判为 BLOCKED 并停车。该记录不能视为完成率。',12,RED,True)

# 8. Method comparison with original TARE
s=slide('与原版 TARE：保留探索框架，改变局部可行性判断','当前 XYθ 版本的算法结构对比；“原版”指原始 TARE + CMU 执行链路')
headers=[('环节','原版 TARE','当前 HT × TARE × XYθ')]
rows=[
 ('探索引导','全局探索 + 局部覆盖视点','保留 TARE 的全局引导与视点采样'),
 ('地形判断','几何地形分析 / 碰撞判断','几何判断 + HT 八方向概率硬门槛'),
 ('局部路线','视点图 / 路径优化；不显式搜索 θ 状态','从里程计 (x,y,θ) 做 Hybrid A*'),
 ('车辆动作','航点由 CMU pathFollower 跟踪','前进 / 倒车 / 原地转向；独立跟踪器'),
]
xs=[.69,3.25,7.49];ws=[2.39,4.06,5.13]
for j,(label,x,w) in enumerate(zip(headers[0],xs,ws)):
    rect(s,x,1.53,w,.56,NAVY,radius=True)
    txt(s,x+.10,1.64,w-.20,.27,label,15,WHITE,True)
for i,row in enumerate(rows):
    y=2.20+i*.89
    for j,(value,x,w) in enumerate(zip(row,xs,ws)):
        rect(s,x,y,w,.77,PALE if i%2==0 else RGBColor(250,252,253),radius=True)
        txt(s,x+.11,y+.16,w-.22,.46,value,14,INK,j==0)
rect(s,.71,6.06,11.90,.59,RGBColor(231,245,244),radius=True)
txt(s,.95,6.18,11.46,.34,'共同点：视点由 TARE 采样；当前版只在局部候选中结合 XYθ 到达时间，尚未做全局联合搜索。',15,INK,True)

# 9. Measured comparison: fresh current-version result only after both trials validate.
if fresh:
    s=slide('最新对照：原版 TARE vs 当前 XYθ 版本',
            f'Garage 各 {comparison["duration_s"]} 秒；相同车辆、64 线仿真雷达、起始姿态与 0.5 m/s 速度上限')
else:
    s=slide('历史对照：Campus 300 秒，原版 TARE vs 早期 HT 集成',
            '相同车辆、64 线仿真雷达、0.5 m/s 速度上限；两次独立运行，HT 组不是当前 XYθ 版本')
cols=[.82,4.44,8.36];widths=[3.35,3.66,3.66]
for x,w,h in zip(cols,widths,['指标','原版 TARE','当前 XYθ' if fresh else '早期 HT + TARE']):
    rect(s,x,1.55,w,.57,NAVY,radius=True)
    txt(s,x+.13,1.68,w-.26,.28,h,16,WHITE,True)
if fresh:
    a=comparison['original'];b=comparison['current']
    def number(d,key,unit='',digits=1):
        v=d.get(key)
        return '无数据' if v is None else f'{v:.{digits}f}{unit}'
    historical=[
      ('行驶距离',number(a,'distance_m',' m',2),number(b,'distance_m',' m',2)),
      ('到访 1 m XY 栅格',str(a['visited_1m_xy_cells'])+' 个',str(b['visited_1m_xy_cells'])+' 个'),
      ('零速指令占比',number({'v':100*a['zero_cmd_fraction']},'v','%',1),
                       number({'v':100*b['zero_cmd_fraction']},'v','%',1)),
      ('倒车指令占比',number({'v':100*a['reverse_cmd_fraction']},'v','%',1),
                       number({'v':100*b['reverse_cmd_fraction']},'v','%',1)),
      ('TARE 规划周期中位数',number({'v':1000*a['planner_runtime_median_s']},'v',' ms',1),
                         number({'v':1000*b['planner_runtime_median_s']},'v',' ms',1)),
      ('收到探索完成信号','是' if a['exploration_finished'] else '否',
                          '是' if b['exploration_finished'] else '否'),
    ]
else:
    historical=[
     ('行驶距离','139.98 m','122.90 m'),
     ('到访 1 m XY 栅格','156 个','91 个'),
     ('零速指令占比','0.3%','8.0%'),
     ('倒车指令占比','5.2%','27.8%'),
     ('TARE 规划周期中位数','393.5 ms','394.5 ms'),
     ('收到探索完成信号','否','否'),
    ]
for i,row in enumerate(historical):
    y=2.20+i*.60
    for j,(value,x,w) in enumerate(zip(row,cols,widths)):
        rect(s,x,y,w,.52,PALE if i%2==0 else RGBColor(250,252,253),radius=True)
        txt(s,x+.13,y+.10,w-.26,.32,value,14,INK,j==0)
rect(s,.82,6.12,11.20,.56,RGBColor(255,240,224),radius=True)
txt(s,1.00,6.24,10.87,.29,
    '到访栅格是路径覆盖代理，非探索体积；单次端到端对照不能证明单一算法的因果收益。' if fresh else
    '到访栅格是路径覆盖代理，非探索体积；单次独立运行不能归因或说明当前 XYθ 性能。',13,RED,True)

# Illustrated diagnosis from the verified October 12 experiments.
D=HERE/'diagnostic_assets'
if (D/'ht_probability_0.png').exists():
    def photo_page(title,kicker,asset,caption):
        page=slide(title,kicker)
        picture(page,D/asset,.62,1.43,12.06,4.96)
        rect(page,.7,6.45,11.94,.49,PALE,radius=True)
        txt(page,.88,6.53,11.58,.32,caption,13,INK,True)
        page.notes_slide.notes_text_frame.text='固定诊断资料：2026-10-12。图像来源：presentations/diagnostic_assets/'+asset+'。'
        return page

    s=slide('改造后的地图链路：全局保留历史，局部完成推理与规划',
            '全局保存高程；每帧按当前位置裁图，HT 模型和 XYθ 搜索仍处理局部窗口')
    stages=[('CuPy 滚动高程图','291 × 291 格\n约 12 m × 12 m'),
            ('全局历史高程图','固定 map 坐标\n高程 / 方差 / 时间'),
            ('局部高程裁图','291 × 291 格\n多帧观测共同支撑'),
            ('八方向 HT 层','241 × 241 × 8\n约 10 m × 10 m')]
    for i,(head,body) in enumerate(stages):
        x=.65+i*3.18
        card(s,x,1.6,2.75,1.9,head,body,TEAL if i%2==0 else BLUE)
        if i<3:txt(s,x+2.78,2.21,.38,.5,'→',23,TEAL,True)
    picture(s,D/'global_history_stop.png',.72,3.75,5.6,2.87)
    card(s,6.65,3.95,5.9,2.42,'历史保存 ≠ 自动判为可通行',
         '未观测地形保持未知；局部图按概率阈值与车体支撑范围做硬拒绝。历史观测默认超过 600 秒后不参与裁图。',ORANGE)
    s.notes_slide.notes_text_frame.text='代码：ht_cupy_bridge_ws/historical_elevation.py。方差加权更新使用受限增益，避免把高度相关的 CuPy 快照当独立测量。全局按稀疏分块存储。'

    photo_page('全局历史高程图：车走过的区域仍留在地图里',
               '14:00 复现实验结束时的全部历史高程；颜色表示高度，白色表示没有有效观测',
               'global_history_stop.png','当前实现保存历史高程分块；完整八方向 HT 仅在局部窗口推理。')
    photo_page('几轮修改后的里程曲线：改善有限，停滞仍存在',
               '统一取各次记录的前 90 秒；同一 Garage、相近初始姿态，探索过程有随机性',
               'distance_90s.png','单次故障定位实验，不能将曲线差异直接解释为某项修改的因果收益。')
    photo_page('车辆实际走到了哪里：轨迹与停滞点',
               '三个代表性运行；绿色圆点为起点，红色叉号为最终停滞位置',
               'trajectories.png','当前版可继续走到更远位置，但尚未完成探索；路径到访格也不是探索体积。')
    photo_page('停滞现场：车辆停止时 HT 地图仍保持更新',
               '14:00 复现运行，共记录 120 秒；下图地图年龄来自独立记录器',
               'stop_timeline.png','本轮未记录到 HT_MAP_STALE；停止主要表现为 NO_EXECUTABLE_ROUTE。')
    photo_page('逐步定位故障：地图过期问题与路线问题',
               '历史融合 → 独立地图回调 → 修正单帧视线前瞻；各次记录时长为 90 或 120 秒',
               'diagnostic_bars.png','独立回调后地图过期状态降为 0，但零速比例仍高，路线与目标选择仍是瓶颈。')

    for first,label in [(0,'0°—135°'),(4,'180°—315°')]:
        photo_page(f'停住时的 HT 原始可通过性：{label}',
                   '同一地形、同一 XY 范围；蓝圈为车辆 0.6 m 支撑半径，箭头为车头方向',
                   f'ht_probability_{first}.png','模型概率：绿色高、红色低、灰色未知；该现场是复现结果，并非 10:10 实验原始快照。')
    for first,label in [(0,'0°—135°'),(4,'180°—315°')]:
        photo_page(f'规划器实际看到的 HT 硬约束：{label}',
                   '概率阈值 0.5；在每个方向上按 0.6 m 支撑半径检查周边格子',
                   f'ht_gate_{first}.png','绿色允许、红色拒绝、灰色未知；来路出现大片红区，上方仍存在绿色通道。')
    environment=HERE/'scheduled_results/stop_ht_2026-10-12_14-00-05/current'
    if (environment/'environment_oblique.png').exists():
        for title,views in [
            ('停滞位置的实际仿真环境：小车与周围地面', [('oblique','斜上方观察：地面、立柱与车辆'),('top','棚内俯视：小车周围的近距离地面')]),
            ('停滞位置的前后视野：墙面与车库通道', [('front','沿停车车头方向观察'),('rear','沿车头反方向观察')])]:
            page=slide(title,'按记录停车姿态在同一静态 Garage 场景还原；不是停车瞬间的历史相机帧')
            for i,(name,label) in enumerate(views):
                picture(page,environment/f'environment_{name}.png',.7+6.1*i,1.68,5.86,4.4)
                txt(page,.82+6.1*i,6.18,5.62,.38,label,15,INK,True)
            txt(page,.82,6.67,11.7,.27,'车辆坐标 (-0.13, 13.71) m，车头约 183.4°；图像由 Gazebo 原生相机渲染。',11,MID)
            page.notes_slide.notes_text_frame.text='还原记录：environment_reconstruction.json；静态场景相机画面，不能单凭图像认定 HT 为误判。'

    s=slide('停滞时的关键区别：车所在点可通行，目标路线仍可能不可达',
            '14:00 现场：位置 (-0.13, 13.71) m，车头约 183.4°；状态 NO_EXECUTABLE_ROUTE')
    picture(s,D/'direction_difference.png',.72,1.47,6.1,5.28)
    card(s,7.0,1.62,5.52,1.37,'起点八方向均允许通过','车所在格子的八方向概率均约 0.999；起点本身没有被 HT 全部拒绝。',TEAL)
    card(s,7.0,3.17,5.52,1.37,'路径连接受硬约束限制','红色区域阻断了部分候选视点路线；绿色通道能否被目标选择利用仍需核对。',ORANGE)
    card(s,7.0,4.72,5.52,1.37,'红色是模型判断','需要将原始高程、真实场景与被拒边一起核对，才能判断真实障碍或模型误判。',RED)

    s=slide('最新复现结果与待验证的根因',
            '2026-10-12 14:00，当前历史地图＋XYθ 版本；Garage 记录 120 秒')
    card(s,.7,1.5,3.82,1.6,'11.11 m','累计行驶距离',BLUE)
    card(s,4.76,1.5,3.82,1.6,'14 个 1 m 格','车辆轨迹到访格',TEAL)
    card(s,8.82,1.5,3.82,1.6,'未完成探索','零速指令占比 75.5%',ORANGE)
    picture(s,D/'ht_gate_0.png',.72,3.42,7.15,2.85)
    txt(s,8.12,3.65,4.48,2.46,'已观察：起点有效，地图新鲜；候选目标存在大量 HT 边拒绝。\n\n待验证：来路低概率是地形事实、融合误差还是模型误判；上方通道为何未被充分利用。',17,INK)
    s.notes_slide.notes_text_frame.text='原始数据：presentations/scheduled_results/stop_ht_2026-10-12_14-00-05/current/summary.json、stopped_ht.npz、tare.log。不同模块根因尚未完全确认。'

# Limitations / next
s=slide('当前边界与下一步','组会结论：闭环路径规划已经跑通；完整 XYθ 探索与实车安全性仍需验证')
card(s,.68,1.50,5.86,1.63,'1  目标选择仍受局部候选限制','已用 XYθ 路线时间评估局部候选收益；全局目标与可达中间目标还需联合处理。',TEAL)
card(s,6.78,1.50,5.86,1.63,'2  停滞尚未解决','地图过期状态已在复测中消除；继续核对 HT 误判、目标连通性与安全回退路径。',ORANGE)
card(s,.68,3.40,5.86,1.63,'3  运动学不等于动力学','现有速度/角速度约束和跟踪限幅，尚未建模滑移、制动、轮矩与坡面牵引。',BLUE)
card(s,6.78,3.40,5.86,1.63,'4  评测范围仍有限','补做相同初始条件的完整任务、多场景和返航测试；再逐步进入实车。',RED)
rect(s,.70,5.68,11.94,.67,NAVY,radius=True)
txt(s,.99,5.82,11.35,.37,'下一里程碑：联合“覆盖收益 + XYθ 可达耗时”，并验证停滞恢复。',19,WHITE,True)

# 11. Backup / sources
s=slide('实验口径与资料','用于组会问答；本页可在正式汇报时略过')
refs=[
 ('系统与代码','HT_README.md；src/tare_planner/include/ht_cost/xytheta_planner.h'),
 ('XYθ 闭环记录','docs/XYTHETA_EXPERIMENT_2026-10-12.md'),
 ('历史运行对比','docs/XYTHETA_COMPARISON_2026-10-12.md'),
 ('可视化原始数据','ht_cupy_bridge_ws/results/xytheta_live_capture_2026-10-12/'),
 ('原版 Campus 对照','tare_planner_upstream_baseline/comparison/results/*/summary.json'),
]
for i,(a,b) in enumerate(refs):
    y=1.38+i*.98
    rect(s,.76,y,11.80,.84,PALE,radius=True)
    txt(s,1.02,y+.15,2.24,.4,a,17,INK,True)
    txt(s,3.38,y+.18,8.87,.35,b,13,MID)
txt(s,.85,6.42,11.45,.35,'注：历史 300 秒合成 HT 测试与当前真实方向概率闭环不同，不在本次定量结果中混用。',12,RED)

OUT.parent.mkdir(parents=True,exist_ok=True)
temporary=OUT.with_name(OUT.name+'.tmp')
prs.save(temporary)
os.replace(temporary,OUT)
shutil.copyfile(OUT,LATEST_OUT)
print(OUT)
